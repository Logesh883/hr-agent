from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from rag.parsing import UnsupportedPolicyFile, to_text


def test_markdown_and_text_pass_through() -> None:
    data = "﻿# Leave Policy\n\n## 1. Entitlements\n".encode()

    assert to_text(data, mime_type="text/markdown; charset=utf-8", file_name="p.md") == (
        "# Leave Policy\n\n## 1. Entitlements\n"
    )
    assert to_text(b"plain", mime_type="application/octet-stream", file_name="p.txt") == "plain"


def test_docx_keeps_headings_lists_and_tables() -> None:
    document = Document()
    document.add_heading("Leave Policy", level=0)  # the Title style
    document.add_heading("1. Entitlements", level=1)
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip(
        [*table.rows[0].cells, *table.rows[1].cells],
        ["Type", "Days", "Annual leave", "18"],
        strict=True,
    ):
        cell.text = value
    document.add_paragraph("Unused leave does not carry over.")
    document.add_paragraph("Nobody approves their own leave.", style="List Bullet")
    buffer = BytesIO()
    document.save(buffer)

    text = to_text(
        buffer.getvalue(),
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_name="leave.docx",
    )

    assert text == (
        "# Leave Policy\n\n"
        "## 1. Entitlements\n\n"
        "| Type | Days |\n| --- | --- |\n| Annual leave | 18 |\n\n"
        "Unused leave does not carry over.\n\n"
        "- Nobody approves their own leave."
    )


def _pdf(text: str) -> bytes:
    """A one-page PDF with a text layer, written by hand (pypdf can't add text)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


def test_pdf_text_layer_is_extracted() -> None:
    text = to_text(_pdf("Late after 10:30"), mime_type="application/pdf", file_name="a.pdf")

    assert text == "Late after 10:30"


def test_a_pdf_without_text_is_refused_not_indexed_empty() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = BytesIO()
    writer.write(buffer)

    with pytest.raises(UnsupportedPolicyFile, match="no text layer"):
        to_text(buffer.getvalue(), mime_type="application/pdf", file_name="scan.pdf")


def test_unknown_types_are_refused() -> None:
    with pytest.raises(UnsupportedPolicyFile):
        to_text(b"\x89PNG", mime_type="image/png", file_name="policy.png")
